"""Export WB declaration workbooks grouped by cabinet and supply; no WB login."""
import argparse
import csv
import json
import os
import re
from collections import defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from wb_declarations import (PROJECT, KEY, RECEIPTS, SHIPMENTS, BASE_TEMPLATE, Sheets, base_vendor,
                             index, integer, normalize, parse_date, load_receipts,
                             shipment_queue, prepare, write_plan)

HEADERS = ['Баркод товара', 'Номер ДТ', 'Артикул', 'Артикул продавца',
           'Наименование товара', 'Количество', 'Код страны происхождения']


def safe_name(value):
    result = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')
    if not result or result.upper().split('.')[0] in {
            'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)),
            *(f'LPT{i}' for i in range(1, 10))}:
        result = '_' + result
    return result[:100]


def read_archive(folder):
    catalog = defaultdict(list)
    warnings = []
    for path in sorted(folder.glob('template*.xlsx')):
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                if len(workbook.worksheets) != 1:
                    continue
                rows = iter(workbook.active.iter_rows(values_only=True))
                headers = next(rows, ())
                columns = {name: index(headers, name) for name in HEADERS}
                for number, row in enumerate(rows, 2):
                    if not any(value is not None for value in row):
                        continue
                    def cell(name):
                        i = columns[name]
                        return row[i] if i < len(row) else None
                    nm = str(integer(cell('Артикул')))
                    barcode = str(cell('Баркод товара') or '').strip()
                    vendor = str(cell('Артикул продавца') or '').strip()
                    if not re.fullmatch(r'\d{8,14}', barcode):
                        raise ValueError(f'Invalid or missing barcode at row {number}')
                    base_vendor(vendor)
                    catalog[nm].append(dict(nm_id=nm, barcode=barcode, vendor_code=vendor,
                                            source=str(path), source_row=number,
                                            cabinet='', supply_id=''))
            finally:
                workbook.close()
        except Exception as exc:
            warnings.append(dict(file=str(path), reason=type(exc).__name__))
    return catalog, warnings


def read_catalog_csv(path):
    # Database export: cabinet;supply_id;nm_id;vendor_code;barcode.
    text = path.read_text(encoding='utf-8-sig')
    dialect = csv.Sniffer().sniff(text[:8192], delimiters=';,\t,')
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    required = {'cabinet', 'supply_id', 'nm_id', 'vendor_code', 'barcode'}
    if not required.issubset(reader.fieldnames or []):
        raise ValueError('Catalog CSV requires cabinet,supply_id,nm_id,vendor_code,barcode')
    result = defaultdict(list)
    for n, row in enumerate(reader, 2):
        nm = str(integer(row['nm_id']))
        barcode = row['barcode'].strip()
        if not re.fullmatch(r'\d{8,14}', barcode):
            raise ValueError(f'Invalid barcode in catalog row {n}')
        base_vendor(row['vendor_code'])
        if not row['cabinet'].strip():
            raise ValueError(f'Missing cabinet in catalog row {n}')
        row.update(nm_id=nm, supply_id=str(integer(row['supply_id'])),
                   source=str(path), source_row=n)
        result[nm].append(row)
    return result


def read_database_catalog(records, barcode_field='barcode'):
    import psycopg
    missing = [name for name in ('PGHOST', 'PGUSER', 'PGPASSWORD', 'PGDATABASE') if not os.environ.get(name)]
    if missing:
        raise ValueError('Для чтения БД добавьте секреты: ' + ', '.join(missing))
    catalog = defaultdict(list)
    with psycopg.connect(host=os.environ['PGHOST'],
                         port=int(os.environ.get('PGPORT', '5432')),
                         user=os.environ['PGUSER'], password=os.environ['PGPASSWORD'],
                         dbname=os.environ['PGDATABASE'], connect_timeout=15,
                         sslmode=os.environ.get('PGSSLMODE', 'prefer'),
                         options='-c default_transaction_read_only=on -c statement_timeout=30000') as connection:
        columns = [row[0] for row in connection.execute(
            '''SELECT attname FROM pg_attribute
               WHERE attrelid = 'wb_supplies_goods'::regclass
               AND attnum > 0 AND NOT attisdropped ORDER BY attnum''')]
        if barcode_field not in columns:
            raise ValueError('Поле баркода не найдено в wb_supplies_goods. Поля таблицы: ' + ', '.join(columns))
        # JSON lookup is parameterized; a configured field name is never interpolated into SQL.
        query = '''SELECT wsg.id, wsg.nm_id, wsg.vendor_code, wsg.quantity,
                          wsg.created_at, to_jsonb(wsg)->>%s AS barcode,
                          a.local_vendor_code, a.account
                   FROM wb_supplies_goods wsg
                   LEFT JOIN article a ON a.nm_id = wsg.nm_id
                   WHERE wsg.id = ANY(%s)
                   ORDER BY wsg.created_at DESC'''
        snapshots = defaultdict(list)
        for supply, nm, vendor, qty, created, barcode, local, account in connection.execute(
                query, (barcode_field, [integer(r['supply_id']) for r in records])):
            if account is None:
                continue
            snapshots[(str(supply), str(nm), normalize(account))].append(dict(
                cabinet=account, supply_id=str(supply), nm_id=str(nm), vendor_code=vendor,
                quantity=qty, created=created, barcode=barcode or '', local=local,
                source='PostgreSQL:wb_supplies_goods', source_row='', origin='DATABASE'))
        for entries in snapshots.values():
            latest = entries[0]['created']
            for entry in entries:
                if entry['created'] == latest:
                    catalog[entry['nm_id']].append(entry)
    return catalog


def identify(record, archive, catalog):
    nm = str(integer(record['nm_id']))
    exact = [item for item in catalog.get(nm, [])
             if normalize(item['cabinet']) == normalize(record['cabinet'])
             and item['supply_id'] == str(record['supply_id'])]
    choices = exact or archive.get(nm, [])
    if not choices:
        raise ValueError('Нет баркода и точного артикула продавца')
    variants = {(str(r['barcode']), r['vendor_code']) for r in choices}
    if len(variants) != 1:
        raise ValueError('Разные баркоды или артикулы для одного артикула WB — нужна проверка')
    barcode, vendor = next(iter(variants))
    if not re.fullmatch(r'\d{8,14}', barcode):
        raise ValueError('В данных поставки отсутствует корректный баркод товара')
    if base_vendor(vendor) != base_vendor(record['wild']):
        raise ValueError('Артикул продавца не соответствует wild из плана')
    if record.get('vendor_code') and record['vendor_code'] != vendor:
        raise ValueError('Артикул из каталога не совпадает с артикулом из БД')
    if exact:
        for entry in exact:
            if 'quantity' in entry and integer(entry['quantity']) != integer(record['quantity']):
                raise ValueError('Количество в БД не совпадает с количеством в плане')
            if entry.get('local') and base_vendor(entry['local']) != base_vendor(record['wild']):
                raise ValueError('Артикул справочника БД не совпадает с wild в плане')
    return dict(barcode=barcode, vendor_code=vendor, nm_id=nm,
                source=choices[-1]['source'], source_row=choices[-1]['source_row'],
                origin=choices[-1].get('origin', 'DB_CSV') if exact else 'ARCHIVED_TEMPLATE')


def make_workbook(base, records, path):
    workbook = load_workbook(base)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Expected one sheet in the WB base template')
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        cols = {h: index(headers, h) + 1 for h in HEADERS}
        styles = {col: copy(sheet.cell(2, col)._style) for col in cols.values()}
        height = sheet.row_dimensions[2].height
        # Discard the original goods; keep the WB header, sheet name and styles.
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        sheet.data_validations.dataValidation = []
        for n, (record, identity) in enumerate(records, 2):
            values = [identity['barcode'], record['dt'], integer(record['nm_id']),
                      identity['vendor_code'], record['product'], integer(record['quantity']),
                      record['country']]
            for header, value in zip(HEADERS, values):
                cell = sheet.cell(n, cols[header], value)
                cell._style = copy(styles[cols[header]])
                if header in ('Баркод товара', 'Номер ДТ', 'Артикул продавца', 'Код страны происхождения'):
                    cell.number_format = '@'
            sheet.row_dimensions[n].height = height
        # No file is written until all rows in this supply have been checked.
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    finally:
        workbook.close()


def export(plan_path, archive_folder, output, catalog_path=None, base_template=None,
           database=False, barcode_field='barcode', use_archive=False):
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    if output.exists():
        raise ValueError('Папка результата уже существует; укажите новую папку')
    archive, warnings = read_archive(archive_folder) if use_archive else ({}, [])
    if database and catalog_path:
        raise ValueError('Выберите один источник данных: БД или CSV')
    catalog = read_database_catalog(plan['records'], barcode_field) if database else (
        read_catalog_csv(catalog_path) if catalog_path else {})
    sources = sorted(archive_folder.glob('template*.xlsx'))
    if base_template is None:
        if sources:
            base_template = sources[0]
        elif BASE_TEMPLATE.is_file():
            base_template = BASE_TEMPLATE
        else:
            raise ValueError('Нет исходного XLSX-шаблона WB; укажите --base-template')
    groups = defaultdict(list)
    for record in plan['records']:
        groups[(record['cabinet'], str(record['supply_id']))].append(record)
    output.mkdir(parents=True)
    log, identities = [], {}
    for (cabinet, supply), records in groups.items():
        status, reason, filename, prepared = 'ГОТОВО', '', '', []
        try:
            if len({str(integer(r['nm_id'])) for r in records}) != len(records):
                raise ValueError('Повтор артикула WB в плане — требуется проверка')
            for record in records:
                if record['state'] != 'READY':
                    raise ValueError(record.get('reason') or 'ДТ не рассчитана')
                identity = identify(record, archive, catalog)
                prepared.append((record, identity))
            if database:
                db_articles = {item['nm_id'] for entries in catalog.values() for item in entries
                               if normalize(item['cabinet']) == normalize(cabinet)
                               and item['supply_id'] == supply}
                if db_articles != {str(integer(r['nm_id'])) for r in records}:
                    raise ValueError('Полный состав поставки в БД не совпадает с планом')
            filename = str(Path(safe_name(cabinet)) / f'{safe_name(supply)}.xlsx')
            make_workbook(base_template, prepared, output / filename)
            identities[f'{cabinet}/{supply}'] = [i for _, i in prepared]
        except ValueError as exc:
            status, reason = 'НЕ СОЗДАН', str(exc)
        log.append(dict(cabinet=cabinet, supply_id=supply, state=status,
                        wild=', '.join(r['wild'] for r in records), filename=filename,
                        reason=reason, source='; '.join(sorted({i['origin'] for _, i in prepared}))))
    with (output / 'СПИСОК_ФАЙЛОВ.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle, delimiter=';')
        writer.writerow(['Кабинет', 'Номер поставки', 'wild', 'Результат', 'Файл', 'Причина', 'Источник артикула и баркода'])
        for row in log:
            writer.writerow([row[k] for k in ('cabinet', 'supply_id', 'wild', 'state', 'filename', 'reason', 'source')])
    (output / 'sources.json').write_text(json.dumps(dict(plan=str(plan_path),
        generated_at=datetime.now(ZoneInfo('Europe/Moscow')).isoformat(), identities=identities,
        warnings=warnings), ensure_ascii=False, indent=2), encoding='utf-8')
    with (output / 'НЕДОСТАЮЩИЕ_ДАННЫЕ.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle, delimiter=';')
        writer.writerow(['cabinet', 'supply_id', 'nm_id', 'vendor_code', 'barcode'])
        failed = {(r['cabinet'], r['supply_id']) for r in log if r['state'] != 'ГОТОВО'}
        for record in plan['records']:
            if (record['cabinet'], str(record['supply_id'])) in failed:
                writer.writerow([record['cabinet'], record['supply_id'], record['nm_id'], '', ''])
    (output / 'ПРОЧИТАЙТЕ.txt').write_text(
        'Файлы сгруппированы по кабинету и номеру поставки. Загружайте каждый XLSX в соответствующую поставку.\n'
        'Вход в WB не использовался. Галочки в Google-таблицах не менялись.\n'
        'ARCHIVED_TEMPLATE: артикул продавца и баркод взяты из прошлых скачанных шаблонов по точному артикулу WB.\n'
        'Перед первой загрузкой сверьте их с текущей поставкой, если карточка товара изменялась.\n'
        'DB_CSV: данные взяты из вашего CSV по кабинету, поставке и артикулу WB.\n'
        'DATABASE: данные взяты напрямую из БД с проверкой количества и состава поставки.\n'
        'Файлы без баркода, точного артикула продавца или рассчитанной ДТ не создаются.\n'
        'НЕДОСТАЮЩИЕ_ДАННЫЕ.csv можно заполнить или заменить экспортом БД и передать как --catalog.\n', encoding='utf-8-sig')
    created = sum(r['state'] == 'ГОТОВО' for r in log)
    return {'supplies': len(log), 'created': created, 'missing': len(log) - created,
            'output': str(output)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group()
    source.add_argument('--plan', type=Path)
    source.add_argument('--date', type=parse_date)
    parser.add_argument('--cabinet', help='Exact cabinet name, case-insensitive')
    parser.add_argument('--supply', help='Supply number')
    parser.add_argument('--key', type=Path, default=KEY)
    parser.add_argument('--templates', type=Path, default=Path.home() / 'Downloads')
    parser.add_argument('--catalog', type=Path)
    parser.add_argument('--with-db', action='store_true')
    parser.add_argument('--db-barcode-field', default='barcode')
    parser.add_argument('--use-archive', action='store_true', help='Explicitly use historical templates as product data')
    parser.add_argument('--base-template', type=Path, default=BASE_TEMPLATE)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    try:
        stamp = datetime.now(ZoneInfo('Europe/Moscow')).strftime('%Y%m%d-%H%M%S-%f')
        plan_path = args.plan
        if plan_path is None:
            sheets = Sheets(args.key)
            as_of = datetime.now(ZoneInfo('Europe/Moscow')).date()
            receipts = load_receipts(sheets.read(RECEIPTS))
            selected, records = shipment_queue(sheets.read(SHIPMENTS), args.date, as_of,
                cabinet=args.cabinet, supply_id=args.supply,
                include_checked=bool(args.cabinet or args.supply))
            prepare(records, receipts, as_of)
            run = PROJECT / 'runs' / stamp
            write_plan(records, selected, as_of, run)
            plan_path = run / 'plan.json'
        payload = json.loads(plan_path.read_text(encoding='utf-8'))
        if args.plan and (args.cabinet or args.supply):
            raise ValueError('Use --cabinet/--supply with live data, not --plan')
        selected = payload.get('shipment_date') or 'no-date'
        output = args.output or PROJECT / 'exports' / f'{selected}_{stamp}'
        result = export(plan_path, args.templates, output, args.catalog, args.base_template,
                        args.with_db, args.db_barcode_field, args.use_archive)
        print(json.dumps(result, ensure_ascii=True))
    except (ValueError, OSError) as exc:
        parser.exit(1, f'{exc}\n')
    except Exception as exc:
        # Authentication failures may include connection details; do not print driver messages.
        parser.exit(1, f'{type(exc).__name__}: ошибка доступа к данным; проверьте секреты и соединение.\n')


if __name__ == '__main__':
    main()
