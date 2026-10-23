import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import telegram_bot as app
from wb_declarations import PROJECT, configured_path
from export_templates import HEADERS


class DistributionTests(unittest.TestCase):
    def test_relative_paths_resolve_against_project(self):
        self.assertEqual(configured_path('secrets/test.json'), PROJECT/'secrets/test.json')
        self.assertEqual(configured_path(str(PROJECT/'assets')), PROJECT/'assets')

    def test_packaged_template_has_only_headers_and_empty_goods(self):
        workbook = load_workbook(PROJECT/'assets/wb_declaration_template.xlsx')
        try:
            self.assertEqual(len(workbook.worksheets), 1)
            sheet = workbook.active
            self.assertEqual([c.value for c in sheet[1]], HEADERS)
            self.assertTrue(all(c.value is None for row in sheet.iter_rows(min_row=2) for c in row))
            self.assertFalse(workbook._external_links)
        finally:
            workbook.close()

    def test_second_instance_is_rejected_and_lock_released(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(app, 'PROJECT', Path(temp)):
            with app.single_instance():
                with self.assertRaises(RuntimeError):
                    with app.single_instance():
                        self.fail('A second process lock was allowed')
            with app.single_instance():
                pass


if __name__ == '__main__':
    unittest.main()
