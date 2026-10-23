import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from wb_declarations import base_vendor, choose_dt, fill_template, integer


class RulesTests(unittest.TestCase):
    def receipts(self):
        return [dict(arrival='20.10.2026', quantity='100', dt='10131010/201026/1234567', row=5),
                dict(arrival='19.10.2026', quantity='100', dt='10131010/191026/1234567', row=6)]

    def test_exact_stock_uses_second_receipt(self):
        dt, trace = choose_dt(self.receipts(), 200, date(2026, 10, 22))
        self.assertEqual(dt, '10131010/191026/1234567')
        self.assertEqual(trace[-1]['cumulative'], 200)

    def test_newest_receipt_can_cover_stock(self):
        self.assertEqual(choose_dt(self.receipts(), 50, date(2026, 10, 22))[0],
                         '10131010/201026/1234567')

    def test_insufficient_stock_and_missing_dt_stop(self):
        with self.assertRaises(ValueError):
            choose_dt(self.receipts(), 201, date(2026, 10, 22))
        receipts = self.receipts()
        receipts[0]['dt'] = ''
        with self.assertRaises(ValueError):
            choose_dt(receipts, 50, date(2026, 10, 22))

    def test_future_receipt_is_excluded(self):
        with self.assertRaises(ValueError):
            choose_dt(self.receipts(), 101, date(2026, 10, 19))

    def test_same_date_keeps_sheet_order_and_uses_quantity(self):
        receipts = self.receipts()
        receipts[1]['arrival'] = receipts[0]['arrival']
        self.assertEqual(choose_dt(receipts, 150, date(2026, 10, 22))[0],
                         '10131010/191026/1234567')
        self.assertEqual(choose_dt(list(reversed(receipts)), 50, date(2026, 10, 22))[0],
                         '10131010/201026/1234567')

    def test_vendor_suffix_and_quantity(self):
        self.assertEqual(base_vendor('wild399D'), 'wild399')
        self.assertEqual(base_vendor('wild399d'), 'wild399')
        self.assertEqual(base_vendor('wild1586d2'), 'wild1586')
        self.assertEqual(base_vendor('wild1863D3'), 'wild1863')
        self.assertEqual(integer('14\u00a0340'), 14340)
        with self.assertRaises(ValueError):
            integer('12.5')

    def test_fill_preserves_exact_vendor_and_original(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            original, output, plan = root/'original.xlsx', root/'filled.xlsx', root/'plan.json'
            wb = Workbook()
            wb.active.append(['Баркод товара', 'Номер ДТ', 'Артикул', 'Артикул продавца',
                              'Наименование товара', 'Количество', 'Код страны происхождения'])
            wb.active.append(['1234567890123', '', 123456789, 'wild1826d', 'Вентиляторы', 42, ''])
            wb.save(original)
            original_bytes = original.read_bytes()
            record = dict(supply_id='12345678', state='READY', cabinet='МАГАЗИНА', nm_id='123456789',
                          wild='wild1826', vendor_code='wild1826d', quantity='42',
                          dt='10131010/201026/1234567', country='156')
            plan.write_text(json.dumps(dict(records=[record])), encoding='utf-8')
            fill_template(plan, original, '12345678', output)
            result = load_workbook(output).active
            self.assertEqual(result['D2'].value, 'wild1826d')
            self.assertEqual(result['A2'].value, '1234567890123')
            self.assertEqual(result['F2'].value, 42)
            self.assertEqual(result['B2'].value, record['dt'])
            self.assertEqual(result['G2'].value, '156')
            self.assertEqual(original.read_bytes(), original_bytes)
            record['quantity'] = '43'
            plan.write_text(json.dumps(dict(records=[record])), encoding='utf-8')
            with self.assertRaises(ValueError):
                fill_template(plan, original, '12345678', root/'wrong.xlsx')
            self.assertFalse((root/'wrong.xlsx').exists())


if __name__ == '__main__':
    unittest.main()
