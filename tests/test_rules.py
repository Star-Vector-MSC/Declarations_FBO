import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from wb_declarations import base_vendor, choose_dt, fill_template, integer, choose_vendor_dt, load_receipts, prepare


class RulesTests(unittest.TestCase):
    def receipts(self):
        return [dict(arrival='20.10.2026', quantity='100', dt='10131010/201026/1234567', row=5),
                dict(arrival='19.10.2026', quantity='100', dt='10131010/191026/1234567', row=6)]

    def test_exact_stock_uses_second_receipt(self):
        dt, trace = choose_dt(self.receipts(), 200, date(2026, 10, 22))
        self.assertEqual(dt, '10131010/191026/1234567')
        self.assertEqual(trace[-1]['cumulative'], 200)

    def test_newest_receipt_can_cover_stock(self):
        self.assertEqual(choose_dt(self.receipts(), 50, date(2026, 10, 22))[0],
                         '10131010/201026/1234567')

    def test_insufficient_stock_and_missing_dt_stop(self):
        with self.assertRaises(ValueError):
            choose_dt(self.receipts(), 201, date(2026, 10, 22))
        receipts = self.receipts()
        receipts[0]['dt'] = ''
        with self.assertRaises(ValueError):
            choose_dt(receipts, 50, date(2026, 10, 22))

    def test_future_receipt_is_excluded(self):
        with self.assertRaises(ValueError):
            choose_dt(self.receipts(), 101, date(2026, 10, 19))

    def test_same_date_keeps_sheet_order_and_uses_quantity(self):
        receipts = self.receipts()
        receipts[1]['arrival'] = receipts[0]['arrival']
        self.assertEqual(choose_dt(receipts, 150, date(2026, 10, 22))[0],
                         '10131010/191026/1234567')
        self.assertEqual(choose_dt(list(reversed(receipts)), 50, date(2026, 10, 22))[0],
                         '10131010/201026/1234567')

    def test_vendor_suffix_and_quantity(self):
        self.assertEqual(base_vendor('wild399D'), 'wild399')
        self.assertEqual(base_vendor('wild399d'), 'wild399')
        self.assertEqual(base_vendor('wild1586d2'), 'wild1586')
        self.assertEqual(base_vendor('wild1863D3'), 'wild1863')
        self.assertEqual(integer('14\u00a0340'), 14340)
        with self.assertRaises(ValueError):
            integer('12.5')

    def test_fill_preserves_exact_vendor_and_original(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            original, output, plan = root/'original.xlsx', root/'filled.xlsx', root/'plan.json'
            wb = Workbook()
            wb.active.append(['Баркод товара', 'Номер ДТ', 'Артикул', 'Артикул продавца',
                              'Наименование товара', 'Количество', 'Код страны происхождения'])
            wb.active.append(['1234567890123', '', 123456789, 'wild1826d', 'Вентиляторы', 42, ''])
            wb.save(original)
            original_bytes = original.read_bytes()
            record = dict(supply_id='12345678', state='READY', cabinet='МАГАЗИНА', nm_id='123456789',
                          wild='wild1826', vendor_code='wild1826d', quantity='42',
                          dt='10131010/201026/1234567', country='156')
            plan.write_text(json.dumps(dict(records=[record])), encoding='utf-8')
            fill_template(plan, original, '12345678', output)
            result = load_workbook(output).active
            self.assertEqual(result['D2'].value, 'wild1826d')
            self.assertEqual(result['A2'].value, '1234567890123')
            self.assertEqual(result['F2'].value, 42)
            self.assertEqual(result['B2'].value, record['dt'])
            self.assertEqual(result['G2'].value, '156')
            self.assertEqual(original.read_bytes(), original_bytes)
            record['quantity'] = '43'
            plan.write_text(json.dumps(dict(records=[record])), encoding='utf-8')
            with self.assertRaises(ValueError):
                fill_template(plan, original, '12345678', root/'wrong.xlsx')
            self.assertFalse((root/'wrong.xlsx').exists())


class KitRulesTests(unittest.TestCase):
    def receipts(self):
        # Two different shipments with matching components; exact equality picks the second.
        return {component: [dict(arrival='20.10.2026', quantity='100',
                    dt='10131010/201026/1234567', row=10+i, vehicle='CAR-NEW',truck='NEW'),
                dict(arrival='19.10.2026', quantity='100',dt='10131010/191026/1234567',
                    row=20+i,vehicle='CAR-OLD',truck='OLD')]
                for i,component in enumerate(('wild1945','wild2135'))}

    def test_components_cover_stock_and_share_boundary_vehicle_and_dt(self):
        dt,trace=choose_vendor_dt('wild103',self.receipts(),200,date(2026,10,22))
        self.assertEqual(dt,'10131010/191026/1234567')
        self.assertEqual({r['component'] for r in trace},{'wild1945','wild2135'})
        self.assertEqual([r['cumulative'] for r in trace],[100,200,100,200])

    def test_equal_dt_with_different_vehicles_is_blocked(self):
        receipts=self.receipts()
        receipts['wild2135'][1]['vehicle']='OTHER'
        with self.assertRaisesRegex(ValueError,'машина'):
            choose_vendor_dt('wild103',receipts,200,date(2026,10,22))

    def test_same_vehicle_with_different_dt_is_blocked(self):
        receipts=self.receipts()
        receipts['wild2135'][1]['dt']='10131010/191026/7654321'
        with self.assertRaisesRegex(ValueError,'ДТ компонентов'):
            choose_vendor_dt('wild103',receipts,200,date(2026,10,22))

    def test_planned_component_without_arrival_does_not_count(self):
        receipts=self.receipts()
        for row in receipts['wild2135']: row['arrival']=''
        with self.assertRaisesRegex(ValueError,'wild2135'):
            choose_vendor_dt('wild103',receipts,200,date(2026,10,22))

    def test_missing_component_does_not_use_direct_kit_receipt(self):
        receipts=self.receipts()
        receipts['wild103']=receipts.pop('wild2135')
        with self.assertRaisesRegex(ValueError,'wild2135'):
            choose_vendor_dt('wild103',receipts,200,date(2026,10,22))

    def test_transport_must_be_present_but_common_truck_can_identify_it(self):
        receipts=self.receipts()
        for component in receipts.values():
            component[1]['vehicle']=''
            component[1]['truck']=''
        with self.assertRaisesRegex(ValueError,'машина'):
            choose_vendor_dt('wild103',receipts,200,date(2026,10,22))
        for component in receipts.values(): component[1]['truck']='TRUCK-2'
        self.assertEqual(choose_vendor_dt('wild103',receipts,200,date(2026,10,22))[0],
                         '10131010/191026/1234567')

    def test_prepare_keeps_finished_kit_article(self):
        records=[dict(wild='wild103D',free_stock='200',quantity='5',cabinet='Shop',product='Комплект')]
        prepare(records,self.receipts(),date(2026,10,22))
        self.assertEqual(records[0]['state'],'READY')
        self.assertEqual(records[0]['wild'],'wild103D')
        self.assertEqual(records[0]['dt'],'10131010/191026/1234567')

    def test_receipts_loader_reads_transport_by_header(self):
        rows=[['wild','Кол-во к заказу','ФАКТИЧЕСКАЯ ДАТА ПРИБЫТИЯ НА СКЛАД ОТ НАС',
               'НОМЕР ТАМОЖЕННОЙ ДЕКЛАРАЦИИ','Номер Трака','Номер ТС'],
              ['wild1945','100','20.10.2026','10131010/201026/1234567','TRUCK','CAR']]
        result=load_receipts(rows)['wild1945'][0]
        self.assertEqual((result['vehicle'],result['truck']),('CAR','TRUCK'))



if __name__ == '__main__':
    unittest.main()
