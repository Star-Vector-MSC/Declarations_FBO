import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from export_templates import HEADERS, export


class ExportTests(unittest.TestCase):
    def test_groups_supplies_preserves_suffix_and_blocks_missing_barcode(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root / 'downloads'
            archive.mkdir()
            source = archive / 'template.xlsx'
            wb = Workbook()
            wb.active.append(HEADERS)
            wb.active.append(['1234567890123', 'old', 123456789, 'wild1826d', 'Old', 5, '003'])
            wb.save(source)
            original = source.read_bytes()
            rows = [dict(cabinet='МАГАЗИНА', supply_id='12345678', nm_id='123456789',
                         wild='wild1826', vendor_code='', quantity='42', product='Вентиляторы',
                         dt='10131010/201026/1234567', country='156', state='READY'),
                    dict(cabinet='Другой кабинет', supply_id='2', nm_id='999', wild='wild999',
                         vendor_code='', quantity='10', product='Казаны',
                         dt='10131010/201026/1234567', country='860', state='READY')]
            plan = root / 'plan.json'
            plan.write_text(json.dumps({'records': rows}), encoding='utf-8')
            result = export(plan, archive, root / 'output', use_archive=True)
            self.assertEqual((result['created'], result['missing']), (1, 1))
            filled = load_workbook(root / 'output' / 'МАГАЗИНА' / '12345678.xlsx')
            values = [c.value for c in filled.active[2]]
            self.assertEqual(values, ['1234567890123', '10131010/201026/1234567',
                                     123456789, 'wild1826d', 'Вентиляторы', 42, '156'])
            filled.close()
            self.assertEqual(source.read_bytes(), original)
            self.assertFalse((root/'output'/'Другой кабинет'/'2.xlsx').exists())

    def test_conflicting_archive_vendor_is_not_guessed(self):
        from export_templates import identify
        record = dict(nm_id='1', cabinet='A', supply_id='2', wild='wild399', vendor_code='')
        archive = {'1': [dict(barcode='12345678', vendor_code='wild399'),
                         dict(barcode='12345678', vendor_code='wild399D')]}
        with self.assertRaises(ValueError):
            identify(record, archive, {})

    def test_catalog_uses_cabinet_supply_and_wb_article(self):
        from export_templates import identify
        record = dict(nm_id='1', cabinet='МАГАЗИНА', supply_id='12345678',
                      wild='wild399', vendor_code='')
        catalog = {'1': [dict(cabinet='Другой', supply_id='12345678', nm_id='1',
                              barcode='11111111', vendor_code='wild399', source='db.csv', source_row=2),
                         dict(cabinet='МАГАЗИНА', supply_id='12345678', nm_id='1',
                              barcode='22222222', vendor_code='wild399d', source='db.csv', source_row=3)]}
        chosen = identify(record, {}, catalog)
        self.assertEqual(chosen['barcode'], '22222222')
        self.assertEqual(chosen['vendor_code'], 'wild399d')

    def test_new_product_is_created_from_catalog_without_previous_shipments(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root / 'downloads'
            archive.mkdir()
            base = archive / 'template.xlsx'
            wb = Workbook()
            wb.active.append(HEADERS)
            wb.active.append(['11111111', '', 111, 'wild111', 'Old product', 1, '156'])
            wb.save(base)
            record = dict(cabinet='Новый кабинет', supply_id='99', nm_id='999', wild='wild999',
                          vendor_code='', quantity='42', product='Казаны',
                          dt='10131010/201026/1234567', country='860', state='READY')
            plan = root / 'plan.json'
            plan.write_text(json.dumps({'records': [record]}), encoding='utf-8')
            catalog = root / 'catalog.csv'
            catalog.write_text('cabinet;supply_id;nm_id;vendor_code;barcode\n'
                               'Новый кабинет;99;999;wild999D;22222222\n', encoding='utf-8-sig')
            result = export(plan, archive, root/'output', catalog_path=catalog)
            self.assertEqual(result['created'], 1)
            result_wb = load_workbook(root/'output'/'Новый кабинет'/'99.xlsx')
            row = [c.value for c in result_wb.active[2]]
            result_wb.close()
            self.assertEqual(row, ['22222222', record['dt'], 999, 'wild999D', 'Казаны', 42, '860'])


if __name__ == '__main__':
    unittest.main()
